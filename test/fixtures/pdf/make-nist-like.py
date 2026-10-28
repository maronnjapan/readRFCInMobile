# NIST SP 風の文書 (Word 由来の 1 段組、Tagged PDF、しおり付き) を作る、テスト用データの生成スクリプト。
#   pip install python-docx   (PDF への変換に LibreOffice を使う)
#   python3 make-nist-like.py nist-like.docx
#   soffice --headless --convert-to 'pdf:writer_pdf_Export:{"UseTaggedPDF":{"type":"boolean","value":"true"},"ExportBookmarks":{"type":"boolean","value":"true"}}' nist-like.docx
#   node ../../../scripts/pdf-extract.mjs nist-like.pdf > nist-like.json
import sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ('left_margin', 'right_margin'):
    setattr(sec, side, Inches(1))
sec.different_first_page_header_footer = True

def page_field(par):
    run = par.add_run()
    for kind, text in (('begin', None), (None, 'PAGE'), ('end', None)):
        if kind:
            el = OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'), kind); run._r.append(el)
        else:
            el = OxmlElement('w:instrText'); el.set(qn('xml:space'), 'preserve'); el.text = text; run._r.append(el)

hp = sec.header.paragraphs[0]
hp.text = 'NIST SP 800-999\tEXAMPLE DIGITAL GUIDELINES'
hp.style.font.size = Pt(9)
fp = sec.footer.paragraphs[0]
fp.text = 'This publication is available free of charge from: https://doi.org/10.6028/NIST.SP.800-999\t'
page_field(fp)

normal = doc.styles['Normal']
normal.font.name = 'Liberation Serif'
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(8)

def para(text, style=None, align=None, size=None, bold=False):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    if size: r.font.size = Pt(size)
    r.bold = bold
    if align: p.alignment = align
    return p

# 表紙
para('NIST Special Publication 800-999', align=WD_ALIGN_PARAGRAPH.LEFT, size=26, bold=True)
para('Example Digital Guidelines', size=20)
para('Authentication and Lifecycle Management', size=16)
para('Jane Q. Example\nJohn R. Sample', size=12)
para('This publication is available free of charge from:\nhttps://doi.org/10.6028/NIST.SP.800-999', size=11)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

para('Abstract', size=14, bold=True)
para('These guidelines provide technical requirements for federal agencies implementing digital identity services. '
     'They are not intended to constrain the development or use of standards outside of this purpose. '
     'This document defines technical requirements for each of the three authenticator assurance levels.')
para('Keywords', size=14, bold=True)
para('authentication; authenticator; credential service provider; digital authentication; passwords.')
para('Table of Contents', size=14, bold=True)
for line in ['1 Purpose ........................................................................................ 3',
             '2 Introduction ................................................................................ 3',
             '   2.1 Notations ............................................................................. 4',
             '3 Authenticator Assurance Levels ....................................................... 4',
             'References ....................................................................................... 6']:
    para(line, size=11)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

doc.add_heading('1 Purpose', level=1)
para('This section is informative. This document and its companion documents provide technical guidelines to agencies '
     'for the implementation of digital authentication. The guidelines are intended to be used by agencies that '
     'provide online services to members of the public as well as to their own employees and contractors.')
para('The terms “SHALL” and “SHALL NOT” indicate requirements to be followed strictly in order to conform to the '
     'publication and from which no deviation is permitted. The terms “SHOULD” and “SHOULD NOT” indicate that among '
     'several possibilities one is recommended as particularly suitable [SP800-63].')
doc.add_heading('2 Introduction', level=1)
para('Authentication is the process of determining the validity of one or more authenticators used to claim a '
     'digital identity. Successful authentication provides reasonable risk-based assurances that the subscriber '
     'accessing the service today is the same as the one who accessed the service previously, as described in Section 3.')
para('The following properties are required of every authenticator:')
for t in ['The authenticator SHALL be resistant to replay attacks, as defined in Table 1.',
          'The verifier SHALL use approved encryption and an authenticated protected channel when requesting '
          'memorized secrets in order to provide resistance to eavesdropping and machine-in-the-middle attacks.',
          'Verifiers SHOULD permit claimants to use “paste” functionality when entering a memorized secret.']:
    para(t, style='List Bullet')
doc.add_heading('2.1 Notations', level=2)
para('This guideline uses the following typographical conventions in text:')
for t in ['Specific terms in CAPITALS represent normative requirements.',
          'Terms in italics represent defined terms in the glossary.']:
    para(t, style='List Number')
doc.add_heading('3 Authenticator Assurance Levels', level=1)
para('In order to satisfy the requirements of a given AAL, a claimant SHALL be authenticated with at least a given '
     'level of strength to be recognized as a subscriber. The result of an authentication process is an identifier '
     'that SHALL be used each time that subscriber authenticates to that relying party.')
cap = para('Table 1. Summary of Requirements by AAL', size=10, bold=True)
table = doc.add_table(rows=4, cols=4)
table.style = 'Table Grid'
rows = [('Requirement', 'AAL1', 'AAL2', 'AAL3'),
        ('Permitted authenticator types', 'Memorized Secret; Look-Up Secret', 'MF OTP Device; MF Crypto Software', 'MF Crypto Device'),
        ('Reauthentication', '30 days', '12 hours, or 30 minutes of inactivity', '12 hours, or 15 minutes of inactivity'),
        ('Replay resistance', 'Not required', 'Required', 'Required')]
for r, row in enumerate(rows):
    for c, text in enumerate(row):
        table.cell(r, c).text = text
para('')
doc.add_heading('3.1 Authenticator Assurance Level 1', level=2)
para('AAL1 provides some assurance that the claimant controls an authenticator bound to the subscriber’s account. '
     'AAL1 requires either single-factor or multi-factor authentication using a wide range of available '
     'authentication technologies. Successful authentication requires that the claimant prove possession and '
     'control of the authenticator through a secure authentication protocol.')
para('Example request sent by the verifier:', size=11)
code = para('POST /authenticate HTTP/1.1\nHost: verifier.example.gov\nContent-Type: application/json\n\n{"otp": "123456"}', size=9)
for r in code.runs: r.font.name = 'Liberation Mono'
doc.add_heading('References', level=1)
para('[FIPS140] National Institute of Standards and Technology (2019) Security Requirements for Cryptographic Modules. '
     '(Department of Commerce, Washington, D.C.), Federal Information Processing Standards Publication (FIPS) 140-3.')
para('[SP800-63] Grassi PA, Garcia ME, Fenton JL (2017) Digital Identity Guidelines. (National Institute of Standards '
     'and Technology, Gaithersburg, MD), NIST Special Publication (SP) 800-63-3.')
doc.save(sys.argv[1])
